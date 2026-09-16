"""
Excel Reader Module for Student Certificate Generator.

Responsible for discovering Excel files, inspecting their columns,
and reading student data with strict preservation of string values (e.g. REG NO).
"""

import os
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
import openpyxl


IGNORE_DIRS = {
    ".git", ".venv", "venv", "env", "__pycache__",
    "output", ".idea", ".vscode", "scratch", ".system_generated"
}

REQUIRED_COLUMNS = ["SI NO", "REG NO", "NAME"]


def normalize_col_name(col_name: Any) -> str:
    """Normalize column header: strip whitespace and uppercase."""
    if col_name is None:
        return ""
    # Replace multiple whitespaces/underscores with single space
    cleaned = " ".join(str(col_name).strip().split())
    # Normalize common variations (e.g. 'REG_NO' -> 'REG NO', 'S.NO' -> 'SI NO')
    cleaned_upper = cleaned.upper().replace("_", " ")
    if cleaned_upper in ("S.NO", "S NO", "SL NO", "SL.NO", "SERIAL NO"):
        return "SI NO"
    if cleaned_upper in ("REG NO", "REGNO", "REGISTRATION NO", "REGISTRATION NUMBER", "REG NUMBER"):
        return "REG NO"
    if cleaned_upper in ("NAME", "STUDENT NAME"):
        return "NAME"
    return cleaned_upper


def find_excel_files(root_dir: str) -> List[str]:
    """
    Search root_dir and its subdirectories for .xlsx files,
    ignoring temporary files and common ignored directories.
    """
    root_path = Path(root_dir).resolve()
    found_files: List[str] = []

    # First search immediate directory
    try:
        for entry in os.scandir(root_path):
            if entry.is_file() and entry.name.lower().endswith(".xlsx") and not entry.name.startswith("~$"):
                found_files.append(str(Path(entry.path).resolve()))
    except OSError:
        pass

    # Then search subdirectories
    for current_root, dirs, files in os.walk(root_path):
        # Modify dirs in-place to skip ignored directories
        dirs[:] = [d for d in dirs if d.lower() not in IGNORE_DIRS and not d.startswith(".")]
        
        # Don't re-add files from the root directory
        if Path(current_root).resolve() == root_path:
            continue

        for f in files:
            if f.lower().endswith(".xlsx") and not f.startswith("~$"):
                found_files.append(str(Path(current_root, f).resolve()))

    return found_files


def inspect_excel_columns(file_path: str) -> Tuple[bool, List[str], Dict[str, int]]:
    """
    Inspect the first sheet's header row in an Excel file.
    Returns:
        (is_valid, missing_columns, column_index_mapping)
    """
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        return False, REQUIRED_COLUMNS.copy(), {}

    try:
        ws = wb.active
        if ws is None:
            return False, REQUIRED_COLUMNS.copy(), {}

        header_row = None
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                header_row = row
                break

        if not header_row:
            return False, REQUIRED_COLUMNS.copy(), {}

        col_map: Dict[str, int] = {}
        for idx, cell in enumerate(header_row):
            norm = normalize_col_name(cell)
            if norm in ("SI NO", "REG NO", "NAME") and norm not in col_map:
                col_map[norm] = idx

        missing = [col for col in REQUIRED_COLUMNS if col not in col_map]
        is_valid = len(missing) == 0
        return is_valid, missing, col_map
    finally:
        wb.close()


def detect_student_excel(root_dir: str, explicit_path: Optional[str] = None) -> str:
    """
    Auto-detect the student Excel file in root_dir or validate explicit_path.
    Raises descriptive RuntimeError/FileNotFoundError if missing or ambiguous.
    """
    if explicit_path:
        resolved = Path(explicit_path)
        if not resolved.is_absolute():
            resolved = Path(root_dir) / resolved
        if not resolved.exists():
            raise FileNotFoundError(f"Specified Excel file not found: {explicit_path}")
        is_valid, missing, _ = inspect_excel_columns(str(resolved))
        if not is_valid:
            raise ValueError(
                f"The Excel file '{resolved.name}' is missing required column(s): {', '.join(missing)}"
            )
        return str(resolved)

    excel_files = find_excel_files(root_dir)
    if not excel_files:
        raise FileNotFoundError(
            f"No Excel (.xlsx) file found in current directory:\n{root_dir}"
        )

    # Inspect each found file to check for required columns
    valid_student_files = []
    for f in excel_files:
        is_valid, missing, _ = inspect_excel_columns(f)
        if is_valid:
            valid_student_files.append(f)

    if not valid_student_files:
        # Check if there's any excel file and report missing columns for clarity
        sample = excel_files[0]
        _, missing, _ = inspect_excel_columns(sample)
        raise ValueError(
            f"Found Excel file '{Path(sample).name}', but it is missing required column(s): {', '.join(missing)}.\n"
            f"Expected columns: {', '.join(REQUIRED_COLUMNS)}"
        )

    if len(valid_student_files) > 1:
        # Check if one of them is in the root directory directly
        root_path = Path(root_dir).resolve()
        root_files = [f for f in valid_student_files if Path(f).parent == root_path]
        if len(root_files) == 1:
            return root_files[0]

        file_list = "\n".join(f" - {f}" for f in valid_student_files)
        raise RuntimeError(
            f"Multiple valid student Excel files found:\n{file_list}\n"
            "Please specify the student Excel file using the --excel argument."
        )

    return valid_student_files[0]


def format_reg_no(val: Any) -> str:
    """
    Safely format REG NO as a string, preventing numeric truncation or .0 floating point.
    Preserves leading zeros (e.g. '00123' stays '00123').
    """
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return str(val).strip()
    return str(val).strip()


def format_si_no(val: Any, default_idx: int) -> str:
    """Safely format SI NO as a string."""
    if val is None or (isinstance(val, str) and val.startswith("=")):
        return str(default_idx)
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    cleaned = str(val).strip()
    return cleaned if cleaned else str(default_idx)


def read_students_from_excel(file_path: str) -> List[Dict[str, Any]]:
    """
    Read students from the Excel file.
    Returns a list of student records:
    [
        {
            "row_num": int,
            "si_no": str,
            "reg_no": str,
            "name": str,
            "is_valid": bool,
            "error": Optional[str]
        },
        ...
    ]
    """
    is_valid, missing, col_map = inspect_excel_columns(file_path)
    if not is_valid:
        raise ValueError(
            f"The Excel file is missing required column(s): {', '.join(missing)}"
        )

    wb = openpyxl.load_workbook(file_path, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            return []

        si_col = col_map["SI NO"]
        reg_col = col_map["REG NO"]
        name_col = col_map["NAME"]

        students: List[Dict[str, Any]] = []
        found_header = False
        student_counter = 1

        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            # Check if row is completely blank
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            # Identify header row
            if not found_header:
                # Check if this row is the header
                norm_vals = [normalize_col_name(c) for c in row if c is not None]
                if "REG NO" in norm_vals and "NAME" in norm_vals:
                    found_header = True
                    continue

            # This is a data row
            raw_si = row[si_col] if si_col < len(row) else None
            raw_reg = row[reg_col] if reg_col < len(row) else None
            raw_name = row[name_col] if name_col < len(row) else None

            # Check if this row is completely empty across the key columns
            if raw_si is None and raw_reg is None and raw_name is None:
                continue

            si_no_str = format_si_no(raw_si, student_counter)
            reg_no_str = format_reg_no(raw_reg)
            name_str = str(raw_name).strip() if raw_name is not None else ""

            # Check validity
            is_row_valid = True
            error_reason = None

            if not reg_no_str and not name_str:
                # Blank student data row, skip
                continue

            if not reg_no_str:
                is_row_valid = False
                error_reason = "REG NO is missing"
            elif not name_str:
                is_row_valid = False
                error_reason = "NAME is missing"

            students.append({
                "row_num": row_idx,
                "si_no": si_no_str,
                "reg_no": reg_no_str,
                "name": name_str,
                "is_valid": is_row_valid,
                "error": error_reason
            })

            if is_row_valid:
                student_counter += 1

        return students
    finally:
        wb.close()
